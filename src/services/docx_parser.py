import io
from docx import Document
from src.core.exceptions import AppError
import structlog

logger = structlog.get_logger()

class DocxParser:
    @staticmethod
    async def extract_docx(file_bytes: bytes) -> dict:
        """
        Extracts text from DOCX bytes including paragraphs and tables using python-docx.
        """
        try:
            doc = Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            table_texts = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = " ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        table_texts.append(row_text)

            full_text = "\n".join(paragraphs + table_texts)
            logger.info("docx_parsed_successfully", total_chars=len(full_text))
            return {
                "text": full_text.strip(),
                "metadata": {"paragraphs_count": len(paragraphs), "tables_count": len(doc.tables)}
            }
        except Exception as e:
            logger.error("docx_parsing_failed", error=str(e))
            raise AppError(f"Failed to parse DOCX document: {str(e)}", status_code=400)
