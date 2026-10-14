import io
import pdfplumber
from pypdf import PdfReader
from docx import Document
from src.core.exceptions import AppError
import structlog

logger = structlog.get_logger()

class ResumeParser:
    @staticmethod
    async def extract_pdf(file_bytes: bytes) -> dict:
        """
        Extracts text and metadata from PDF bytes using pdfplumber and pypdf fallback.
        """
        text_content = ""
        metadata = {}
        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                metadata = pdf.metadata or {}
                pages_text = [page.extract_text() or "" for page in pdf.pages]
                text_content = "\n".join(pages_text)

            if not text_content.strip():
                reader = PdfReader(io.BytesIO(file_bytes))
                metadata = reader.metadata or metadata
                pages_text = [page.extract_text() or "" for page in reader.pages]
                text_content = "\n".join(pages_text)

            logger.info("pdf_parsed_successfully", total_chars=len(text_content))
            return {
                "text": text_content.strip(),
                "metadata": {k: str(v) for k, v in metadata.items() if v}
            }
        except Exception as e:
            logger.error("pdf_parsing_failed", error=str(e))
            raise AppError(f"Failed to parse PDF document: {str(e)}", status_code=400)

    @staticmethod
    async def extract_docx(file_bytes: bytes) -> dict:
        """
        Extracts text from DOCX bytes including paragraphs and tables using python-docx.
        """
        try:
            doc = Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            table_texts = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = " ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        table_texts.append(row_text)

            full_text = "\n".join(paragraphs + table_texts)
            logger.info("docx_parsed_successfully", total_chars=len(full_text))
            return {
                "text": full_text.strip(),
                "metadata": {"paragraphs_count": len(paragraphs), "tables_count": len(doc.tables)}
            }
        except Exception as e:
            logger.error("docx_parsing_failed", error=str(e))
            raise AppError(f"Failed to parse DOCX document: {str(e)}", status_code=400)

    @classmethod
    async def parse_document(cls, file_type: str, file_bytes: bytes) -> dict:
        if file_type == "pdf":
            return await cls.extract_pdf(file_bytes)
        elif file_type == "docx":
            return await cls.extract_docx(file_bytes)
        else:
            raise AppError(f"Unsupported document type: {file_type}", status_code=400)
