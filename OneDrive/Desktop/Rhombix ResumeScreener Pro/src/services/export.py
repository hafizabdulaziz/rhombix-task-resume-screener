import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import structlog

logger = structlog.get_logger()

class ReportExporter:
    @staticmethod
    def generate_excel_report(job_title: str, results: list) -> bytes:
        """
        Generates a styled multi-sheet Excel evaluation report for candidate rankings.
        """
        wb = Workbook()
        
        # Sheet 1: Summary Dashboard
        ws_summary = wb.active
        ws_summary.title = "Executive Summary"
        ws_summary.views.sheetView[0].showGridLines = True

        # Styles
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid") # Dark Slate
        accent_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid") # Emerald Green
        title_font = Font(name="Calibri", size=16, bold=True, color="0F172A")
        border_thin = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        # Title Block
        ws_summary["A1"] = f"Candidate Screening Evaluation Report: {job_title}"
        ws_summary["A1"].font = title_font
        ws_summary.append([]) # Empty row

        # Summary Headers
        summary_headers = ["Rank", "Candidate Name", "Email", "Composite Score (%)", "Skill Match (%)", "Experience Score (%)", "Semantic Score (%)", "Status"]
        ws_summary.append(summary_headers)

        for col_idx in range(1, len(summary_headers) + 1):
            cell = ws_summary.cell(row=3, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for idx, res in enumerate(results, start=1):
            row_data = [
                idx,
                res.get("candidate_name", "N/A"),
                res.get("candidate_email", "N/A"),
                res.get("composite_score", 0.0),
                res.get("skill_match_score", 0.0),
                res.get("experience_match_score", 0.0),
                res.get("semantic_score", 0.0),
                res.get("status", "Pending")
            ]
            ws_summary.append(row_data)
            row_num = ws_summary.max_row
            for c_idx in range(1, len(row_data) + 1):
                c = ws_summary.cell(row=row_num, column=c_idx)
                c.border = border_thin
                c.alignment = Alignment(horizontal="left" if c_idx in [2, 3, 8] else "center", vertical="center")
                if c_idx == 4: # Composite score highlight
                    c.font = Font(bold=True, color="047857")

        # Sheet 2: Detailed Skill Breakdown
        ws_details = wb.create_sheet(title="Skill Breakdown Details")
        ws_details.views.sheetView[0].showGridLines = True
        
        detail_headers = ["Candidate Name", "Skill Name", "Matched / Missing", "Match Type"]
        ws_details.append(detail_headers)
        for col_idx in range(1, len(detail_headers) + 1):
            cell = ws_details.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for res in results:
            cand_name = res.get("candidate_name", "N/A")
            for skill_score in res.get("skill_scores", []):
                # Handle both dict or object
                if isinstance(skill_score, dict):
                    s_name = skill_score.get("skill_name")
                    is_match = skill_score.get("is_matched")
                    m_type = skill_score.get("match_type")
                else:
                    s_name = skill_score.skill_name
                    is_match = skill_score.is_matched
                    m_type = skill_score.match_type

                ws_details.append([
                    cand_name,
                    s_name,
                    "Matched" if is_match else "Missing",
                    m_type
                ])
                r_num = ws_details.max_row
                for c_idx in range(1, 5):
                    c = ws_details.cell(row=r_num, column=c_idx)
                    c.border = border_thin
                    c.alignment = Alignment(horizontal="center", vertical="center")

        # Auto-fit column widths
        for ws in [ws_summary, ws_details]:
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    @staticmethod
    def generate_csv_report(results: list) -> str:
        """
        Generates a lightweight CSV summary report string.
        """
        flat_data = []
        for idx, res in enumerate(results, start=1):
            flat_data.append({
                "Rank": idx,
                "Candidate Name": res.get("candidate_name"),
                "Email": res.get("candidate_email"),
                "Composite Score (%)": res.get("composite_score"),
                "Skill Match (%)": res.get("skill_match_score"),
                "Experience Score (%)": res.get("experience_match_score"),
                "Semantic Score (%)": res.get("semantic_score")
            })
        df = pd.DataFrame(flat_data)
        return df.to_csv(index=False)
